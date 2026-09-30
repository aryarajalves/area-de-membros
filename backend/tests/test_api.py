import io
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.main import app, init_superadmin
from app.core.database import Base, get_db
from app.core.config import settings
from app.core.security import get_password_hash
from datetime import datetime, timezone
from app.models.user import User, Invite, RegistrationVerification
from app.models.course import Course, UserCourse, Module, Lesson, LessonProgress
from app.models.webhook import Webhook, WebhookLog

SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()

app.dependency_overrides[get_db] = override_get_db

@pytest.fixture(autouse=True)
def setup_database():
    Base.metadata.create_all(bind=engine)
    # Seed superadmin
    db = TestingSessionLocal()
    superadmin = User(
        email=settings.SUPERADMIN_EMAIL,
        name=settings.SUPERADMIN_NAME,
        hashed_password=get_password_hash(settings.SUPERADMIN_PASSWORD),
        role="superadmin",
        is_active=True
    )
    db.add(superadmin)
    db.commit()
    db.close()
    yield
    Base.metadata.drop_all(bind=engine)

client = TestClient(app)

def get_superadmin_headers():
    res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    return {"Authorization": f"Bearer {res.json()['access_token']}"}

def test_health_and_root():
    res = client.get("/health")
    assert res.status_code == 200
    assert res.json()["status"] == "ok"

    res_root = client.get("/")
    assert res_root.status_code == 200
    assert "Area de Membros" in res_root.json()["message"]

def test_login_superadmin_success():
    login_payload = {
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    }
    res = client.post("/api/v1/auth/login", json=login_payload)
    assert res.status_code == 200
    data = res.json()
    assert "access_token" in data
    assert data["user"]["role"] == "superadmin"

def test_login_wrong_password():
    login_payload = {
        "email": settings.SUPERADMIN_EMAIL,
        "password": "WrongPassword123!"
    }
    res = client.post("/api/v1/auth/login", json=login_payload)
    assert res.status_code == 401

def test_create_invite_and_register_user():
    # Login superadmin to get token
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Tentar criar convite para superadmin (deve falhar)
    invalid_invite = client.post("/api/v1/auth/invites", json={"role": "superadmin", "duration_hours": 12}, headers=headers)
    assert invalid_invite.status_code == 422 or invalid_invite.status_code == 400

    # 2. Criar convite válido para admin
    invite_res = client.post("/api/v1/auth/invites", json={"role": "admin", "duration_hours": 24}, headers=headers)
    assert invite_res.status_code == 201
    invite_data = invite_res.json()
    invite_token = invite_data["token"]
    assert invite_token is not None

    # 3. Validar convite
    val_res = client.get(f"/api/v1/auth/invites/validate?token={invite_token}")
    assert val_res.status_code == 200
    assert val_res.json()["valid"] is True
    assert val_res.json()["role"] == "admin"

    # 4. Registrar usuário com senha fraca (deve falhar por validação de senha)
    weak_register = client.post("/api/v1/auth/register", json={
        "token": invite_token,
        "name": "Novo Admin",
        "email": "admin@teste.com",
        "password": "fraca"
    })
    assert weak_register.status_code == 422

    # 4.1 Registrar com senhas divergentes (deve falhar com 400)
    mismatch_reg = client.post("/api/v1/auth/register", json={
        "token": invite_token,
        "name": "Novo Admin",
        "email": "admin@teste.com",
        "password": "Password123456!",
        "password_confirm": "OutraSenha123456!"
    })
    assert mismatch_reg.status_code == 400
    assert "não coincidem" in mismatch_reg.json()["detail"]

    # 5. Iniciar registro de usuário com senha válida e confirmação coincidente (dispara OTP)
    valid_register = client.post("/api/v1/auth/register", json={
        "token": invite_token,
        "name": "Novo Admin",
        "email": "admin@teste.com",
        "password": "Password123456!",
        "password_confirm": "Password123456!"
    })
    assert valid_register.status_code == 200
    assert valid_register.json()["requires_verification"] is True

    # 5.1 Obter código OTP gerado no banco e validar código incorreto
    db = TestingSessionLocal()
    verif = db.query(RegistrationVerification).filter(RegistrationVerification.email == "admin@teste.com").first()
    assert verif is not None
    otp_code = verif.code
    db.close()

    wrong_code_res = client.post("/api/v1/auth/register/verify", json={
        "email": "admin@teste.com",
        "code": "000000"
    })
    assert wrong_code_res.status_code == 400
    assert "inválido" in wrong_code_res.json()["detail"]

    # 5.2 Validar código correto e concluir criação da conta
    verify_res = client.post("/api/v1/auth/register/verify", json={
        "email": "admin@teste.com",
        "code": otp_code
    })
    assert verify_res.status_code == 201
    new_user = verify_res.json()
    assert new_user["email"] == "admin@teste.com"
    assert new_user["role"] == "admin"

    # 6. Tentar usar o mesmo convite novamente (deve falhar pois já foi usado)
    reuse_invite = client.post("/api/v1/auth/register", json={
        "token": invite_token,
        "name": "Outro Admin",
        "email": "outro@teste.com",
        "password": "Password123456!"
    })
    assert reuse_invite.status_code == 400

    # 7. Listar convites gerados
    list_inv = client.get("/api/v1/auth/invites", headers=headers)
    assert list_inv.status_code == 200
    inv_list_data = list_inv.json()
    assert len(inv_list_data) >= 1
    found_invite = next(i for i in inv_list_data if i["token"] == invite_token)
    assert found_invite["is_used"] is True
    assert found_invite["used_by_email"] == "admin@teste.com"
    assert "restantes" in found_invite["time_remaining"] or "Expirado" in found_invite["time_remaining"]

def test_duplicate_email_registration_forbidden():
    # Login superadmin
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Criar convite
    invite_res = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 24}, headers=headers)
    invite_token = invite_res.json()["token"]

    # Tentar registrar com email que já é do superadmin
    dup_res = client.post("/api/v1/auth/register", json={
        "token": invite_token,
        "name": "Tentativa Clone",
        "email": settings.SUPERADMIN_EMAIL,
        "password": "SecurePassword123!"
    })
    assert dup_res.status_code == 400
    assert "Já existe um usuário" in dup_res.json()["detail"]

def test_edit_user_role_and_superadmin_protection():
    # Login superadmin
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Tentar editar o próprio Super Admin (deve falhar com 403)
    edit_superadmin = client.patch("/api/v1/auth/users/1", json={"role": "user"}, headers=headers)
    assert edit_superadmin.status_code == 403

    # 2. Criar um usuário via convite
    inv_res = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 10}, headers=headers)
    inv_tok = inv_res.json()["token"]
    reg_res = client.post("/api/v1/auth/register", json={
        "token": inv_tok,
        "name": "Membro Equipe",
        "email": "membro@equipe.com",
        "password": "MinhaSenhaForte1@",
        "password_confirm": "MinhaSenhaForte1@"
    })
    assert reg_res.status_code == 200

    db = TestingSessionLocal()
    verif = db.query(RegistrationVerification).filter(RegistrationVerification.email == "membro@equipe.com").first()
    otp_code = verif.code
    db.close()

    verify_res = client.post("/api/v1/auth/register/verify", json={
        "email": "membro@equipe.com",
        "code": otp_code
    })
    assert verify_res.status_code == 201
    user_id = verify_res.json()["id"]

    # 3. Alterar perfil do usuário para 'admin'
    edit_res = client.patch(f"/api/v1/auth/users/{user_id}", json={"role": "admin"}, headers=headers)
    assert edit_res.status_code == 200
    assert edit_res.json()["role"] == "admin"

    # 4. Alterar de volta para 'user'
    edit_back = client.patch(f"/api/v1/auth/users/{user_id}", json={"role": "user"}, headers=headers)
    assert edit_back.status_code == 200
    assert edit_back.json()["role"] == "user"

def test_password_reset_flow():
    # Login superadmin
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Tentar gerar link de redefinição para o Super Admin (deve ser proibido com 403)
    reset_sa = client.post("/api/v1/auth/users/1/reset-password-request", headers=headers)
    assert reset_sa.status_code == 403

    # 2. Criar um usuário comum
    inv_res = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 10}, headers=headers)
    inv_tok = inv_res.json()["token"]
    reg_res = client.post("/api/v1/auth/register", json={
        "token": inv_tok,
        "name": "Usuario Teste",
        "email": "userteste@equipe.com",
        "password": "SenhaAntiga12345!",
        "password_confirm": "SenhaAntiga12345!"
    })
    assert reg_res.status_code == 200

    db = TestingSessionLocal()
    verif = db.query(RegistrationVerification).filter(RegistrationVerification.email == "userteste@equipe.com").first()
    otp_code = verif.code
    db.close()

    verify_res = client.post("/api/v1/auth/register/verify", json={
        "email": "userteste@equipe.com",
        "code": otp_code
    })
    assert verify_res.status_code == 201
    user_id = verify_res.json()["id"]

    # 3. Gerar link de redefinição de senha para esse usuário
    reset_req = client.post(f"/api/v1/auth/users/{user_id}/reset-password-request", headers=headers)
    assert reset_req.status_code == 200
    reset_data = reset_req.json()
    reset_token = reset_data["token"]
    assert reset_data["reset_url"].startswith("/reset-password?token=")

    # 4. Validar token de redefinição
    val_res = client.get(f"/api/v1/auth/reset-password/validate?token={reset_token}")
    assert val_res.status_code == 200
    assert val_res.json()["valid"] is True
    assert val_res.json()["email"] == "userteste@equipe.com"

    # 5. Tentar redefinir com senhas divergentes (deve falhar com 400)
    mismatch_res = client.post("/api/v1/auth/reset-password/confirm", json={
        "token": reset_token,
        "password": "NovaSenhaSegura123@",
        "password_confirm": "SenhaTotalmenteDiferente1@"
    })
    assert mismatch_res.status_code == 400
    assert "não coincidem" in mismatch_res.json()["detail"]

    # 6. Confirmar redefinição com nova senha coincidente
    confirm_res = client.post("/api/v1/auth/reset-password/confirm", json={
        "token": reset_token,
        "password": "NovaSenhaSegura123@",
        "password_confirm": "NovaSenhaSegura123@"
    })
    assert confirm_res.status_code == 200

    # 7. Testar login com a nova senha
    login_new = client.post("/api/v1/auth/login", json={
        "email": "userteste@equipe.com",
        "password": "NovaSenhaSegura123@"
    })
    assert login_new.status_code == 200

def test_argon2id_hash_and_pepper():
    from app.core.security import get_password_hash, verify_password
    plain = "MinhaSenhaSuperSecreta123!"
    hashed = get_password_hash(plain)
    
    # Verifica que o hash utiliza Argon2id ($argon2id$)
    assert hashed.startswith("$argon2id$")
    # Verifica que o hash é verificado com sucesso com a senha correta
    assert verify_password(plain, hashed) is True
    # Verifica que uma senha incorreta falha
    assert verify_password("SenhaErrada123!", hashed) is False

def test_user_and_invite_deletion_and_bulk_delete():
    # Login superadmin
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Criar 2 convites
    inv1 = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 10}, headers=headers).json()
    inv2 = client.post("/api/v1/auth/invites", json={"role": "admin", "duration_hours": 10}, headers=headers).json()

    # 2. Excluir convite individual
    del_inv = client.delete(f"/api/v1/auth/invites/{inv1['id']}", headers=headers)
    assert del_inv.status_code == 200

    # 3. Exclusão em massa de convites
    bulk_inv = client.post("/api/v1/auth/invites/bulk-delete", json={"ids": [inv2["id"]]}, headers=headers)
    assert bulk_inv.status_code == 200
    assert bulk_inv.json()["deleted_count"] == 1

    # 4. Criar um usuário para testar deleção
    inv3 = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 10}, headers=headers).json()
    reg = client.post("/api/v1/auth/register", json={
        "token": inv3["token"],
        "name": "Deletar Usuario",
        "email": "deletar@teste.com",
        "password": "SenhaForte123!@",
        "password_confirm": "SenhaForte123!@"
    })
    assert reg.status_code == 200
    db = TestingSessionLocal()
    code = db.query(RegistrationVerification).filter(RegistrationVerification.email == "deletar@teste.com").first().code
    db.close()
    verif = client.post("/api/v1/auth/register/verify", json={"email": "deletar@teste.com", "code": code})
    user_id = verif.json()["id"]

    # 5. Tentar excluir Super Admin (proibido 403)
    del_sa = client.delete("/api/v1/auth/users/1", headers=headers)
    assert del_sa.status_code == 403

    # 6. Excluir usuário individual
    del_u = client.delete(f"/api/v1/auth/users/{user_id}", headers=headers)
    assert del_u.status_code == 200

    # 7. Testar bulk delete de usuários
    bulk_u = client.post("/api/v1/auth/users/bulk-delete", json={"ids": [1]}, headers=headers)
    assert bulk_u.status_code == 200
    # Superadmin é ignorado pelo bulk delete
    assert bulk_u.json()["deleted_count"] == 0

def test_non_superadmin_forbidden_from_management_endpoints():
    # 1. Login superadmin para criar usuário comum
    sa_login = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    sa_token = sa_login.json()["access_token"]
    sa_headers = {"Authorization": f"Bearer {sa_token}"}

    inv = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 2}, headers=sa_headers).json()
    client.post("/api/v1/auth/register", json={
        "token": inv["token"],
        "name": "Comum Acesso",
        "email": "comum.acesso@teste.com",
        "password": "SenhaForte123!@",
        "password_confirm": "SenhaForte123!@"
    })
    db = TestingSessionLocal()
    code = db.query(RegistrationVerification).filter(RegistrationVerification.email == "comum.acesso@teste.com").first().code
    db.close()
    client.post("/api/v1/auth/register/verify", json={"email": "comum.acesso@teste.com", "code": code})

    # 2. Login como usuário comum
    user_login = client.post("/api/v1/auth/login", json={
        "email": "comum.acesso@teste.com",
        "password": "SenhaForte123!@"
    })
    user_token = user_login.json()["access_token"]
    user_headers = {"Authorization": f"Bearer {user_token}"}

    # 3. Usuário comum NÃO pode listar usuários (403)
    users_res = client.get("/api/v1/auth/users", headers=user_headers)
    assert users_res.status_code == 403

    # 4. Usuário comum NÃO pode listar nem criar convites (403)
    inv_list_res = client.get("/api/v1/auth/invites", headers=user_headers)
    assert inv_list_res.status_code == 403

    inv_create_res = client.post("/api/v1/auth/invites", json={"role": "user", "duration_hours": 1}, headers=user_headers)
    assert inv_create_res.status_code == 403

def test_backup_system_full_flow():
    # Login superadmin
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Obter estatísticas do painel
    stats_res = client.get("/api/v1/backups/dashboard-stats", headers=headers)
    assert stats_res.status_code == 200
    assert "current_count" in stats_res.json()
    assert "retention_max" in stats_res.json()

    # 2. Configurar agendador
    sched_res = client.put("/api/v1/backups/schedule", json={
        "is_active": True,
        "frequency": "12h",
        "destination_folder": "projetobase/backups/",
        "retention_max": 25
    }, headers=headers)
    assert sched_res.status_code == 200
    assert sched_res.json()["frequency"] == "12h"
    assert sched_res.json()["retention_max"] == 25

    # 3. Disparar backup manual com nome customizado
    from zoneinfo import ZoneInfo
    from datetime import datetime
    b_now = datetime.now(ZoneInfo("America/Sao_Paulo"))

    manual_res = client.post("/api/v1/backups/manual", json={"custom_name": "teste_brasilia"}, headers=headers)
    assert manual_res.status_code == 201
    bkp = manual_res.json()
    assert bkp["backup_type"] == "manual"
    assert bkp["filename"].startswith("teste_brasilia_")
    assert bkp["filename"].endswith(".dump.gz")
    # Confirma que contém o ano e mês do horário de Brasília
    assert b_now.strftime("%Y_%m") in bkp["filename"]
    backup_id = bkp["id"]

    # 3.1 Renomear backup existente
    rename_res = client.patch(f"/api/v1/backups/{backup_id}/rename", json={"new_name": "backup_renomeado_ok"}, headers=headers)
    assert rename_res.status_code == 200
    assert rename_res.json()["filename"] == "backup_renomeado_ok.dump.gz"

    # 4. Listar backups
    list_res = client.get("/api/v1/backups/list", headers=headers)
    assert list_res.status_code == 200
    assert len(list_res.json()) >= 1

    # 5. Download do backup
    down_res = client.get(f"/api/v1/backups/download/{backup_id}", headers=headers)
    assert down_res.status_code == 200
    assert len(down_res.content) > 0

    # 6. Restauração do backup
    restore_res = client.post(f"/api/v1/backups/restore/{backup_id}", headers=headers)
    assert restore_res.status_code == 200
    assert "restaurado com sucesso" in restore_res.json()["message"]

    # 7. Importar backup externo
    import io
    fake_dump_file = io.BytesIO(b"-- Fake pg dump sql\nCREATE TABLE test_import();")
    import_res = client.post(
        "/api/v1/backups/import",
        files={"file": ("externo_backup.dump.gz", fake_dump_file, "application/gzip")},
        headers=headers
    )
    assert import_res.status_code == 201
    imported_bkp = import_res.json()
    assert imported_bkp["backup_type"] == "imported"

    # 8. Exclusão individual
    del_res = client.delete(f"/api/v1/backups/{backup_id}", headers=headers)
    assert del_res.status_code == 200

    # 9. Exclusão em lote
    bulk_res = client.post("/api/v1/backups/bulk-delete", json={"ids": [imported_bkp["id"]]}, headers=headers)
    assert bulk_res.status_code == 200
    assert bulk_res.json()["deleted_count"] == 1

def test_docker_logs_endpoints():
    login_res = client.post("/api/v1/auth/login", json={
        "email": settings.SUPERADMIN_EMAIL,
        "password": settings.SUPERADMIN_PASSWORD
    })
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Listar serviços disponíveis (apenas backend e frontend)
    services_res = client.get("/api/v1/logs/services", headers=headers)
    assert services_res.status_code == 200
    services = services_res.json()
    assert len(services) == 2
    service_ids = [s["id"] for s in services]
    assert "backend" in service_ids
    assert "frontend" in service_ids
    assert "db" not in service_ids

    # 1.1 Testar conversão de Horário de Brasília
    from app.services.docker_logs_service import format_to_brasilia_time
    nginx_sample = '172.31.0.1 - - [29/Sep/2026:17:59:35 +0000] "GET / HTTP/1.1" 200 123'
    converted = format_to_brasilia_time(nginx_sample)
    # 17:59:35 UTC deve ser convertido para 14:59:35 em Brasília (UTC-3)
    assert "14:59:35" in converted
    assert "29/09/2026 14:59:35" in converted

    # 1.2 Testar conversão de Docker timestamp no início da linha
    docker_sample = '2026-09-29T18:04:29.123456789Z INFO: Servidor iniciado com sucesso.'
    docker_converted = format_to_brasilia_time(docker_sample)
    assert "[29/09/2026 15:04:29]" in docker_converted
    assert "INFO: Servidor iniciado com sucesso." in docker_converted

    # 2. Consultar logs de backend com filtro de data e horário
    logs_res = client.get("/api/v1/logs/backend?tail=20&date=2026-09-29&start_time=14:00&end_time=16:00", headers=headers)
    assert logs_res.status_code == 200
    data = logs_res.json()
    assert data["service"] == "backend"
    assert "logs" in data
    assert isinstance(data["logs"], list)

    # 3. Acesso bloqueado sem token
    unauth_res = client.get("/api/v1/logs/backend")
    assert unauth_res.status_code == 401


def test_export_students_csv_and_xlsx():
    """Valida exportação de alunos em formato CSV e XLSX."""
    headers = get_superadmin_headers()
    db = TestingSessionLocal()
    student = User(
        name="Aluno Exportacao",
        email="aluno_export@teste.com",
        role="aluno",
        hashed_password=get_password_hash("123456"),
        is_active=True
    )
    db.add(student)
    db.commit()
    db.close()

    res_csv = client.get("/api/v1/students/export?format=csv", headers=headers)
    assert res_csv.status_code == 200
    assert "text/csv" in res_csv.headers["content-type"]
    assert "aluno_export@teste.com" in res_csv.text

    res_xlsx = client.get("/api/v1/students/export?format=xlsx", headers=headers)
    assert res_xlsx.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res_xlsx.headers["content-type"]
    assert len(res_xlsx.content) > 0


def test_download_import_template():
    """Valida download de modelo de planilha para importação."""
    headers = get_superadmin_headers()
    res_csv = client.get("/api/v1/students/import/template?format=csv", headers=headers)
    assert res_csv.status_code == 200
    assert "Nome;Email;Senha;Cursos;Tempo de Acesso" in res_csv.text

    res_xlsx = client.get("/api/v1/students/import/template?format=xlsx", headers=headers)
    assert res_xlsx.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
    assert "Modelo_Importacao_Alunos" in wb.sheetnames


def test_import_students_csv_and_xlsx_flow():
    """Valida importação de alunos via CSV e XLSX com cursos e prazos."""
    headers = get_superadmin_headers()
    db = TestingSessionLocal()
    curso = Course(title="Curso Import Flow", description="Desc", is_published=True)
    db.add(curso)
    db.commit()
    course_id = curso.id
    db.close()

    # 1. Importação via CSV
    csv_content = (
        "Nome;Email;Senha;Cursos;Tempo de Acesso\n"
        "Aluno Import CSV;import_csv@teste.com;123456;Curso Import Flow;1 ano\n"
    ).encode("utf-8-sig")

    files = {"file": ("alunos.csv", io.BytesIO(csv_content), "text/csv")}
    res_csv = client.post("/api/v1/students/import", headers=headers, files=files, data={"default_access_duration": "lifetime"})
    assert res_csv.status_code == 200
    payload = res_csv.json()
    assert payload["created_count"] == 1

    db = TestingSessionLocal()
    aluno_csv = db.query(User).filter(User.email == "import_csv@teste.com").first()
    assert aluno_csv is not None
    uc = db.query(UserCourse).filter(UserCourse.user_id == aluno_csv.id).first()
    assert uc is not None
    assert uc.access_duration == "1_year"
    db.close()

    # 2. Importação via XLSX
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alunos"
    ws.append(["Nome", "Email", "Senha", "Cursos", "Tempo de Acesso"])
    ws.append(["Aluno Import XLSX", "import_xlsx@teste.com", "senha123", str(course_id), "6 meses"])
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    files_xlsx = {"file": ("alunos.xlsx", buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res_xlsx = client.post("/api/v1/students/import", headers=headers, files=files_xlsx, data={"default_course_ids": str(course_id)})
    assert res_xlsx.status_code == 200
    assert res_xlsx.json()["created_count"] == 1

    # 3. Atualização (reimportar mesmo email)
    buffer.seek(0)
    files_xlsx2 = {"file": ("alunos.xlsx", buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res_reimport = client.post("/api/v1/students/import", headers=headers, files=files_xlsx2, data={})
    assert res_reimport.status_code == 200
    assert res_reimport.json()["created_count"] == 0
    assert res_reimport.json()["updated_count"] == 1


def test_webhook_integrations_crud_and_validation():
    """Valida criação, listagem, atualização e exclusão de integrações de webhooks."""
    headers = get_superadmin_headers()

    # 1. Listar inicialmente (vazio ou existente)
    res_list = client.get("/api/v1/integrations", headers=headers)
    assert res_list.status_code == 200

    # 2. Criar integração válida
    payload = {
        "name": "Webhook Notificação Alunos",
        "url": "https://api.meusite.com/webhook",
        "events": ["course.progress.25", "course.progress.50", "course.progress.75", "course.progress.100"],
        "secret_key": "chave_secreta_teste",
        "is_active": True
    }
    res_create = client.post("/api/v1/integrations", json=payload, headers=headers)
    assert res_create.status_code == 201
    wh_data = res_create.json()
    assert wh_data["name"] == "Webhook Notificação Alunos"
    assert wh_data["url"] == "https://api.meusite.com/webhook"
    assert "course.progress.50" in wh_data["events"]
    wh_id = wh_data["id"]

    # 3. Consultar detalhes por ID
    res_get = client.get(f"/api/v1/integrations/{wh_id}", headers=headers)
    assert res_get.status_code == 200
    assert res_get.json()["id"] == wh_id

    # 4. Atualizar integração
    res_update = client.put(f"/api/v1/integrations/{wh_id}", json={
        "name": "Webhook Notificação Editado",
        "is_active": False
    }, headers=headers)
    assert res_update.status_code == 200
    assert res_update.json()["name"] == "Webhook Notificação Editado"
    assert res_update.json()["is_active"] is False

    # 5. Acesso não autorizado sem permissão de admin
    db = TestingSessionLocal()
    aluno_user = User(
        name="Aluno Comum",
        email="aluno_comum_wh@teste.com",
        role="aluno",
        hashed_password=get_password_hash("123456"),
        is_active=True
    )
    db.add(aluno_user)
    db.commit()
    db.close()

    aluno_token = client.post("/api/v1/auth/login", json={"email": "aluno_comum_wh@teste.com", "password": "123456"}).json()["access_token"]
    assert client.get("/api/v1/integrations", headers={"Authorization": f"Bearer {aluno_token}"}).status_code == 403

    # 6. Excluir integração
    res_delete = client.delete(f"/api/v1/integrations/{wh_id}", headers=headers)
    assert res_delete.status_code == 200
    assert client.get(f"/api/v1/integrations/{wh_id}", headers=headers).status_code == 404


def test_webhook_test_dispatch_and_logs():
    """Valida disparo de evento de teste de webhook e consulta de logs."""
    headers = get_superadmin_headers()

    # Cria webhook de teste
    res_create = client.post("/api/v1/integrations", json={
        "name": "Webhook de Teste Logs",
        "url": "https://httpbin.org/status/200",
        "events": ["course.progress.100"],
        "is_active": True
    }, headers=headers)
    wh_id = res_create.json()["id"]

    # Dispara teste
    res_test = client.post(f"/api/v1/integrations/{wh_id}/test", headers=headers)
    assert res_test.status_code == 200
    test_json = res_test.json()
    assert "status_code" in test_json

    # Consulta logs do webhook
    res_logs = client.get(f"/api/v1/integrations/{wh_id}/logs", headers=headers)
    assert res_logs.status_code == 200
    logs = res_logs.json()
    assert len(logs) >= 1
    assert logs[0]["event"] == "webhook.test"
    assert "student" in logs[0]["payload"]["data"]

    # Limpeza
    client.delete(f"/api/v1/integrations/{wh_id}", headers=headers)


def test_student_course_history_and_expiration_timeline():
    """Valida cálculo de tempo restante, progresso de validade e histórico detalhado de aulas do aluno."""
    headers = get_superadmin_headers()
    db = TestingSessionLocal()

    # Cria curso com módulo e aulas
    curso = Course(title="Curso Histórico Teste", description="Desc")
    db.add(curso)
    db.commit()

    modulo = Module(title="Módulo 1 - Fundamentos", course_id=curso.id, order_index=1)
    db.add(modulo)
    db.commit()

    aula1 = Lesson(title="Aula 1 - Boas Vindas", module_id=modulo.id, order_index=1)
    aula2 = Lesson(title="Aula 2 - Primeiros Passos", module_id=modulo.id, order_index=2)
    db.add_all([aula1, aula2])
    db.commit()

    aula1_id = aula1.id
    course_id = curso.id

    # Cria aluno
    aluno = User(
        name="Aluno Com Histórico",
        email="aluno_historico@teste.com",
        role="aluno",
        hashed_password=get_password_hash("123456"),
        is_active=True
    )
    db.add(aluno)
    db.commit()

    # Vincula curso com validade de 30 dias a partir de agora
    now = datetime.now(timezone.utc)
    from datetime import timedelta
    expires_at = now + timedelta(days=30)
    uc = UserCourse(
        user_id=aluno.id,
        course_id=course_id,
        access_duration="1_month",
        expires_at=expires_at
    )
    db.add(uc)

    # Marca aula 1 como concluída
    lp = LessonProgress(
        lesson_id=aula1_id,
        user_id=aluno.id,
        is_completed=True,
        completed_at=now
    )
    db.add(lp)
    db.commit()
    db.close()

    # 1. Valida listagem de estudantes com dados de tempo restante
    res_students = client.get("/api/v1/students?search=aluno_historico@teste.com", headers=headers)
    assert res_students.status_code == 200
    st_data = res_students.json()["items"][0]
    assert len(st_data["courses"]) == 1
    c_info = st_data["courses"][0]
    assert c_info["days_remaining"] is not None
    assert c_info["days_remaining"] >= 29
    assert c_info["is_expired"] is False
    assert c_info["time_progress_percent"] is not None

    # 2. Valida consulta do histórico de aulas assistidas
    res_history = client.get(f"/api/v1/students/{st_data['id']}/courses/{c_info['course_id']}/history", headers=headers)
    assert res_history.status_code == 200
    history = res_history.json()
    assert len(history) == 1
    assert history[0]["lesson_id"] == aula1_id
    assert history[0]["lesson_title"] == "Aula 1 - Boas Vindas"
    assert history[0]["module_title"] == "Módulo 1 - Fundamentos"
    assert history[0]["is_completed"] is True
    assert history[0]["completed_at"] is not None


def test_check_renewals_and_events():
    """Valida endpoint POST /api/v1/integrations/check-renewals e eventos de aviso de renovação e expiração."""
    headers = get_superadmin_headers()
    from datetime import timedelta
    db = TestingSessionLocal()

    curso = Course(title="Curso Renovação Teste", description="Desc")
    db.add(curso)
    db.commit()

    aluno_expirado = User(
        name="Aluno Vencido",
        email="vencido@teste.com",
        role="aluno",
        hashed_password=get_password_hash("123456"),
        is_active=True
    )
    aluno_aviso = User(
        name="Aluno Aviso 7d",
        email="aviso7d@teste.com",
        role="aluno",
        hashed_password=get_password_hash("123456"),
        is_active=True
    )
    db.add_all([aluno_expirado, aluno_aviso])
    db.commit()

    now = datetime.now(timezone.utc)
    # 1 expirado
    uc1 = UserCourse(user_id=aluno_expirado.id, course_id=curso.id, expires_at=now - timedelta(days=2))
    # 1 a vencer em 7 dias
    uc2 = UserCourse(user_id=aluno_aviso.id, course_id=curso.id, expires_at=now + timedelta(days=7))
    db.add_all([uc1, uc2])
    db.commit()
    db.close()

    # Cria webhook de escuta para renovação
    res_wh = client.post("/api/v1/integrations", json={
        "name": "Webhook Renovação",
        "url": "https://httpbin.org/status/200",
        "events": ["course.renewal.warning_7d", "course.renewal.expired"],
        "is_active": True
    }, headers=headers)
    assert res_wh.status_code == 201
    wh_id = res_wh.json()["id"]

    # Dispara checagem de renovação
    res_check = client.post("/api/v1/integrations/check-renewals", headers=headers)
    assert res_check.status_code == 200
    check_json = res_check.json()
    assert check_json["success"] is True
    assert "details" in check_json

    # Limpeza
    client.delete(f"/api/v1/integrations/{wh_id}", headers=headers)





