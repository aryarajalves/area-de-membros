import io
import csv
import secrets
import string
from datetime import datetime, timezone
from typing import List, Dict, Any, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from app.api.v1.endpoints.users import calculate_course_expiration

DURATION_MAP = {
    "vitalicio": "lifetime",
    "vitalício": "lifetime",
    "lifetime": "lifetime",
    "1 mes": "1_month",
    "1 mês": "1_month",
    "1_month": "1_month",
    "3 meses": "3_months",
    "3_months": "3_months",
    "6 meses": "6_months",
    "6_months": "6_months",
    "1 ano": "1_year",
    "1_year": "1_year",
    "2 anos": "2_years",
    "2_years": "2_years",
    "3 anos": "3_years",
    "3_years": "3_years",
}

def normalize_duration(duration_str: Optional[str], default: str = "lifetime") -> str:
    if not duration_str:
        return default
    clean = duration_str.strip().lower()
    return DURATION_MAP.get(clean, default)

def generate_random_password(length: int = 12) -> str:
    """Gera uma senha forte atendendo aos requisitos mínimos."""
    alphabet = string.ascii_letters + string.digits + "!@#$%"
    while True:
        pwd = "".join(secrets.choice(alphabet) for _ in range(length))
        if (any(c.islower() for c in pwd)
            and any(c.isupper() for c in pwd)
            and any(c.isdigit() for c in pwd)
            and any(c in "!@#$%" for c in pwd)):
            return pwd

def export_students_csv(students_list: List[Dict[str, Any]]) -> str:
    """Gera o conteúdo de um arquivo CSV com codificação UTF-8."""
    output = io.StringIO()
    # Adiciona BOM para correta visualização no Excel
    output.write('\ufeff')
    writer = csv.writer(output, delimiter=';')
    writer.writerow([
        "ID",
        "Nome",
        "Email",
        "Status",
        "Data de Cadastro",
        "Cursos Liberados",
        "Prazos de Acesso",
        "Progresso Geral (%)",
        "Detalhes do Progresso"
    ])

    for s in students_list:
        courses = s.get("courses", [])
        course_names = ", ".join([c.get("course_title", "") for c in courses])
        
        durations = []
        progress_details = []
        for c in courses:
            dur = c.get("access_duration", "lifetime")
            exp = c.get("expires_at")
            exp_str = f"Expira em {exp.strftime('%d/%m/%Y')}" if exp else "Vitalício"
            durations.append(f"{c.get('course_title')}: {exp_str}")
            progress_details.append(f"{c.get('course_title')}: {c.get('completed_lessons', 0)}/{c.get('total_lessons', 0)} ({c.get('progress_percent', 0)}%)")

        created_at = s.get("created_at")
        created_str = created_at.strftime('%d/%m/%Y %H:%M') if created_at else ""

        writer.writerow([
            s.get("id"),
            s.get("name"),
            s.get("email"),
            "Ativo" if s.get("is_active") else "Inativo",
            created_str,
            course_names,
            " | ".join(durations),
            f"{s.get('overall_progress_percent', 0)}%",
            " | ".join(progress_details)
        ])

    return output.getvalue()

def export_students_xlsx(students_list: List[Dict[str, Any]]) -> bytes:
    """Gera um arquivo Excel binário (.xlsx) estilizado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alunos"

    headers = [
        "ID",
        "Nome",
        "Email",
        "Status",
        "Data de Cadastro",
        "Cursos Liberados",
        "Prazos de Acesso",
        "Progresso Geral (%)",
        "Detalhes do Progresso"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for s in students_list:
        courses = s.get("courses", [])
        course_names = ", ".join([c.get("course_title", "") for c in courses])
        
        durations = []
        progress_details = []
        for c in courses:
            exp = c.get("expires_at")
            exp_str = f"Expira em {exp.strftime('%d/%m/%Y')}" if exp else "Vitalício"
            durations.append(f"{c.get('course_title')}: {exp_str}")
            progress_details.append(f"{c.get('course_title')}: {c.get('completed_lessons', 0)}/{c.get('total_lessons', 0)} ({c.get('progress_percent', 0)}%)")

        created_at = s.get("created_at")
        created_str = created_at.strftime('%d/%m/%Y %H:%M') if created_at else ""

        ws.append([
            s.get("id"),
            s.get("name"),
            s.get("email"),
            "Ativo" if s.get("is_active") else "Inativo",
            created_str,
            course_names,
            " | ".join(durations),
            f"{s.get('overall_progress_percent', 0)}%",
            " | ".join(progress_details)
        ])

    # Ajusta largura das colunas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_template_csv() -> str:
    """Gera um arquivo de exemplo/modelo CSV para importação."""
    output = io.StringIO()
    output.write('\ufeff')
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Nome", "Email", "Senha", "Cursos", "Tempo de Acesso"])
    writer.writerow([
        "Carlos Silva",
        "carlos@exemplo.com",
        "SenhaForte123!",
        "Bussola Astrologica",
        "Vitalício"
    ])
    writer.writerow([
        "Mariana Souza",
        "mariana@exemplo.com",
        "",
        "Bussola Astrologica, Curso 2",
        "1 ano"
    ])
    return output.getvalue()

def generate_template_xlsx() -> bytes:
    """Gera um arquivo de exemplo/modelo Excel para importação."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Modelo_Importacao_Alunos"

    headers = ["Nome", "Email", "Senha", "Cursos", "Tempo de Acesso"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill

    ws.append([
        "Carlos Silva",
        "carlos@exemplo.com",
        "SenhaForte123!",
        "Bussola Astrologica",
        "Vitalício"
    ])
    ws.append([
        "Mariana Souza",
        "mariana@exemplo.com",
        "",
        "Bussola Astrologica, Curso 2",
        "1 ano"
    ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 18)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def parse_imported_file(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Lê o arquivo CSV ou XLSX enviado e retorna uma lista de dicionários normalizados:
    [{'name': ..., 'email': ..., 'password': ..., 'courses': ..., 'access_duration': ...}]
    """
    records = []
    fname_lower = filename.lower()

    if fname_lower.endswith('.xlsx') or fname_lower.endswith('.xls'):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        
        headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        
        # Mapeia colunas
        col_indices = _map_column_indices(headers)

        for row in rows[1:]:
            if not row or all(v is None or str(v).strip() == "" for v in row):
                continue
            record = _extract_row_record(row, col_indices)
            if record.get("email"):
                records.append(record)

    else:
        # Tratamento de CSV com suporte a múltiplos encodings e delimitadores
        text_content = ""
        for encoding in ['utf-8-sig', 'utf-8', 'latin1', 'cp1252']:
            try:
                text_content = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if not text_content:
            text_content = file_bytes.decode('utf-8', errors='ignore')

        # Detecta delimitador (, ou ;)
        sample = text_content[:2048]
        delimiter = ';' if sample.count(';') > sample.count(',') else ','

        reader = csv.reader(io.StringIO(text_content), delimiter=delimiter)
        rows = list(reader)
        if not rows:
            return []

        headers = [h.strip().lower() for h in rows[0]]
        col_indices = _map_column_indices(headers)

        for row in rows[1:]:
            if not row or all(str(v).strip() == "" for v in row):
                continue
            record = _extract_row_record(row, col_indices)
            if record.get("email"):
                records.append(record)

    return records

def _map_column_indices(headers: List[str]) -> Dict[str, int]:
    indices = {"name": -1, "email": -1, "password": -1, "courses": -1, "access_duration": -1}
    for idx, h in enumerate(headers):
        clean_h = h.lower().replace("_", " ").replace("-", " ").strip()
        if clean_h in ["nome", "name", "nome completo", "aluno"]:
            indices["name"] = idx
        elif clean_h in ["email", "e mail", "e-mail"]:
            indices["email"] = idx
        elif clean_h in ["senha", "password", "pass"]:
            indices["password"] = idx
        elif clean_h in ["cursos", "curso", "courses", "produtos"]:
            indices["courses"] = idx
        elif clean_h in ["tempo de acesso", "tempo acesso", "acesso", "validade", "prazo", "duration"]:
            indices["access_duration"] = idx
    return indices

def _extract_row_record(row: Tuple, col_indices: Dict[str, int]) -> Dict[str, Any]:
    def get_val(key: str) -> str:
        idx = col_indices.get(key, -1)
        if idx >= 0 and idx < len(row) and row[idx] is not None:
            return str(row[idx]).strip()
        return ""

    return {
        "name": get_val("name"),
        "email": get_val("email"),
        "password": get_val("password"),
        "courses": get_val("courses"),
        "access_duration": get_val("access_duration")
    }
